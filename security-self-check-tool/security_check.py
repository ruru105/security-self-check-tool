# -*- coding: utf-8 -*-
"""
セキュリティ公開前チェックツール

AIにコードを書いてもらった(バイブコーディングした)あと、公開・納品する前に
「最低限ここだけは確認しよう」というセルフチェックを行うツールです。

裏側では bandit / pip-audit / detect-secrets という、信頼されているOSSの
静的解析ツールを使って自動チェックを行い、その結果を日本語でまとめた
Excelレポートに出力します。あわせて、自動では判断しきれない項目を含む
100項目の手動チェックリストも同じレポートに収録します。

このツールの結論はあくまで「気づきのきっかけ」です。
最終判断はご自身で行ってください。誤検知・見落としもありえます。

使い方:
    python security_check.py <検査したいプロジェクトのフォルダ>

出力:
    security_report.xlsx (検査結果サマリー / 自動チェック結果 / 手動チェックリスト の3シート)
"""

import argparse
import json
import subprocess
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from checklist_data import CHECKLIST, CATEGORIES

# ─────────────────────────────────────────────
# 色・スタイル定義(work1のCSVチェッカーと統一感を持たせる)
# ─────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SEV_FILL = {
    "critical": PatternFill("solid", fgColor="F8696B"),  # 赤系
    "warning": PatternFill("solid", fgColor="FFEB84"),   # 黄系
    "info": PatternFill("solid", fgColor="BDD7EE"),      # 水色系
}
SEV_LABEL_JA = {
    "critical": "重大",
    "warning": "警告",
    "info": "参考",
}


def run_bandit(target: str) -> list:
    """bandit を実行し、Python コードの危険なパターンを検出する。"""
    # テストコード(test_*.py等)はassert多用などでノイズが多いため対象から除外する
    exclude = "*/test_*.py,*test_*.py,*/*_test.py,*_test.py,*/tests/*,*/conftest.py"
    try:
        proc = subprocess.run(
            ["bandit", "-r", target, "-f", "json", "-q", "-x", exclude],
            capture_output=True, text=True, timeout=300,
        )
        # banditは問題を検出すると exit code 1 を返すため、stdoutのJSONを優先して読む
        data = json.loads(proc.stdout) if proc.stdout.strip() else {"results": []}
    except FileNotFoundError:
        print("警告: bandit が見つかりません。pip install bandit でインストールしてください。")
        return []
    except (json.JSONDecodeError, subprocess.TimeoutExpired):
        return []

    findings = []
    sev_map = {"HIGH": "critical", "MEDIUM": "warning", "LOW": "info"}
    for r in data.get("results", []):
        findings.append({
            "ツール": "bandit",
            "重要度": sev_map.get(r.get("issue_severity", "LOW"), "info"),
            "ファイル": r.get("filename", ""),
            "行": r.get("line_number", ""),
            "内容": r.get("issue_text", ""),
            "ルールID": r.get("test_id", ""),
        })
    return findings


def run_pip_audit(target: str) -> list:
    """pip-audit を実行し、依存ライブラリの既知脆弱性(CVE)を検出する。"""
    req_file = None
    for name in ("requirements.txt", "requirements-prod.txt"):
        candidate = Path(target) / name
        if candidate.exists():
            req_file = candidate
            break
    if req_file is None:
        return []

    try:
        proc = subprocess.run(
            ["pip-audit", "-r", str(req_file), "-f", "json"],
            capture_output=True, text=True, timeout=300,
        )
        data = json.loads(proc.stdout) if proc.stdout.strip() else {"dependencies": []}
    except FileNotFoundError:
        print("警告: pip-audit が見つかりません。pip install pip-audit でインストールしてください。")
        return []
    except (json.JSONDecodeError, subprocess.TimeoutExpired):
        return []

    findings = []
    seen = set()  # pip-audit は同じIDの脆弱性情報が複数DBから重複して返ることがあるため、重複を除く
    for dep in data.get("dependencies", []):
        for vuln in dep.get("vulns", []):
            key = (dep["name"], dep["version"], vuln.get("id", ""))
            if key in seen:
                continue
            seen.add(key)
            findings.append({
                "ツール": "pip-audit",
                "重要度": "critical",
                "ファイル": req_file.name,
                "行": "",
                "内容": f"{dep['name']} {dep['version']} に既知の脆弱性 {vuln.get('id', '')}",
                "ルールID": vuln.get("id", ""),
            })
    return findings


def run_detect_secrets(target: str) -> list:
    """detect-secrets を実行し、APIキー等のハードコードを検出する。"""
    try:
        proc = subprocess.run(
            # --all-files を付けないと、gitで管理されていないファイルがスキップされてしまう
            ["detect-secrets", "scan", "--all-files", target],
            capture_output=True, text=True, timeout=300,
        )
        data = json.loads(proc.stdout) if proc.stdout.strip() else {"results": {}}
    except FileNotFoundError:
        print("警告: detect-secrets が見つかりません。pip install detect-secrets でインストールしてください。")
        return []
    except (json.JSONDecodeError, subprocess.TimeoutExpired):
        return []

    findings = []
    for filename, items in data.get("results", {}).items():
        for item in items:
            findings.append({
                "ツール": "detect-secrets",
                "重要度": "warning",
                "ファイル": filename,
                "行": item.get("line_number", ""),
                "内容": f"秘密情報らしき値を検出({item.get('type', '')})。誤検知の可能性もあります",
                "ルールID": item.get("type", ""),
            })
    return findings


def build_report(target: str, out_path: str):
    print(f"検査対象: {target}")
    print("bandit (コードの危険なパターン) を実行中...")
    findings = run_bandit(target)
    print(f"  → {len(findings)} 件")

    print("pip-audit (依存ライブラリの既知脆弱性) を実行中...")
    pa = run_pip_audit(target)
    findings += pa
    print(f"  → {len(pa)} 件")

    print("detect-secrets (秘密情報の直書き) を実行中...")
    ds = run_detect_secrets(target)
    findings += ds
    print(f"  → {len(ds)} 件")

    wb = Workbook()
    _write_summary_sheet(wb, target, findings)
    _write_findings_sheet(wb, findings)
    _write_checklist_sheet(wb)
    wb.save(out_path)
    print(f"\nレポートを作成しました: {out_path}")


def _write_summary_sheet(wb, target, findings):
    ws = wb.active
    ws.title = "検査結果サマリー"
    ws.append(["セキュリティ公開前チェック結果"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["検査対象", target])
    ws.append(["検査日時", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws.append([])

    counts = {"critical": 0, "warning": 0, "info": 0}
    for f in findings:
        counts[f["重要度"]] = counts.get(f["重要度"], 0) + 1

    ws.append(["重要度", "件数"])
    for cell in ws[ws.max_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for sev in ("critical", "warning", "info"):
        ws.append([SEV_LABEL_JA[sev], counts[sev]])
        ws.cell(row=ws.max_row, column=1).fill = SEV_FILL[sev]

    ws.append([])
    ws.append(["総合判定",
               "要確認(自動チェックで検出あり)" if sum(counts.values()) else "自動チェックでは検出なし"])
    ws.append([])
    ws.append(["※ このツールの結論はあくまで「自動チェックでの気づき」です。"])
    ws.append(["　 最終判断はご自身で行ってください。誤検知・見落としの可能性があります。"])
    ws.append(["　 自動チェックの対象は Python コードと requirements.txt のみです。"])

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 50


def _write_findings_sheet(wb, findings):
    ws = wb.create_sheet("自動チェック結果")
    headers = ["重要度", "ツール", "ファイル", "行", "内容", "ルールID"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    if not findings:
        ws.append(["—", "—", "—", "—", "自動チェックでは指摘事項は見つかりませんでした", "—"])
    else:
        # 重要度順(critical→warning→info)に並べる
        order = {"critical": 0, "warning": 1, "info": 2}
        for f in sorted(findings, key=lambda x: order.get(x["重要度"], 9)):
            row = [SEV_LABEL_JA[f["重要度"]], f["ツール"], f["ファイル"], f["行"], f["内容"], f["ルールID"]]
            ws.append(row)
            ws.cell(row=ws.max_row, column=1).fill = SEV_FILL[f["重要度"]]

    widths = [10, 14, 30, 8, 60, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _write_checklist_sheet(wb):
    ws = wb.create_sheet("手動チェックリスト(100項目)")
    headers = ["No.", "カテゴリ", "チェック項目", "確認(OK/NG/-)", "対応のヒント", "自動チェック対象の目安"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    current_cat = None
    for item in CHECKLIST:
        cat_label = CATEGORIES[item["category"]]
        ws.append([
            item["id"],
            cat_label,
            item["item"],
            "",
            item["hint"],
            "○" if item["auto"] else "",
        ])
        row = ws.max_row
        if item["category"] != current_cat:
            # カテゴリの先頭行を少し強調
            for c in range(1, 7):
                ws.cell(row=row, column=c).font = Font(bold=True)
            current_cat = item["category"]

    widths = [6, 32, 55, 14, 45, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    parser = argparse.ArgumentParser(description="セキュリティ公開前チェックツール")
    parser.add_argument("target", help="検査したいプロジェクトのフォルダ")
    parser.add_argument("-o", "--output", default="security_report.xlsx", help="出力するExcelファイル名")
    args = parser.parse_args()

    if not Path(args.target).is_dir():
        print(f"エラー: フォルダが見つかりません: {args.target}")
        sys.exit(1)

    build_report(args.target, args.output)


if __name__ == "__main__":
    main()
