# -*- coding: utf-8 -*-
"""最低限の自動テスト。チェックリストの整合性と、レポート生成が壊れていないかを確認する。"""
from pathlib import Path

from checklist_data import CATEGORIES, CHECKLIST
from security_check import build_report


def test_checklist_has_100_items():
    assert len(CHECKLIST) == 100


def test_checklist_ids_are_unique_and_sequential():
    ids = sorted(item["id"] for item in CHECKLIST)
    assert ids == list(range(1, 101))


def test_checklist_categories_are_all_known():
    for item in CHECKLIST:
        assert item["category"] in CATEGORIES


def test_build_report_on_sample_project(tmp_path):
    out = tmp_path / "report.xlsx"
    build_report("sample_project", str(out))
    assert out.exists()

    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb.sheetnames == ["検査結果サマリー", "自動チェック結果", "手動チェックリスト(100項目)"]
    # サンプルプロジェクトには意図的な問題を仕込んであるので、何かしら検出されるはず
    findings_ws = wb["自動チェック結果"]
    assert findings_ws.max_row > 1
