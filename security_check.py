# -*- coding: utf-8 -*-
"""
セキュリティ公開前チェックツール

AIにコードを書いてもらった(バイブコーディングした)あと、公開・納品する前に
「最低限ここだけは確認しよう」というセルフチェックを行うツールです。

裏側では bandit / pip-audit / detect-secrets という、信頼されているOSSの
静的解析ツールを使って自動チェックを行い、その結果を日本語でまとめた
Excelレポートに出力します。あわせて、自動では判断しきれない項目を含む
100項目の手動チェックリストも同じレポートに収録します。

このツールの結論はあくまで「気づきのきっかけ」です。
最終判断はご自身で行ってください。誤検知・見落としもありえます。

使い方:
    python security_check.py <検査したいプロジェクトのフォルダ>

出力:
    security_report.xlsx (検査結果サマリー / 自動チェック結果 / 手動チェックリスト の3シート)
"""

import argparse
import json
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from checklist_data import CHECKLIST, CATEGORIES

# ─────────────────────────────────────────────
# 色・スタイル定義(work1のCSVチェッカーと統一感を持たせる)
# ─────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SEV_FILL = {
    "critical": PatternFill("solid", fgColor="F8696B"),  # 赤系
    "warning": PatternFill("solid", fgColor="FFEB84"),   # 黄系
    "info": PatternFill("solid", fgColor="BDD7EE"),      # 水色系
}
SEV_LABEL_JA = {
    "critical": "重大",
    "warning": "警告",
    "info": "参考",
}


def _read_scan_result(proc, tool, key, expected_type, allowed_codes):
    """終了コードと結果の形を確認する。空・不正な出力を検出ゼロに置き換えない。"""
    if proc.returncode not in allowed_codes:
        return None, f"{tool} が異常終了しました(終了コード {proc.returncode})"
    if not proc.stdout.strip():
        return None, f"{tool} の検査結果が空です(検査の完了を確認できません)"
    try:
        data = json.loads(proc.stdout)
    except json.JSONDecodeError:
        return None, f"{tool} の出力を読み取れませんでした(JSON形式ではありません)"
    if not isinstance(data, dict) or not isinstance(data.get(key), expected_type):
        return None, f"{tool} の出力に必要な検査結果がありません({key}の形式が不正)"
    entries = data[key]
    if expected_type is list:
        valid_entries = all(isinstance(item, dict) for item in entries)
    else:
        valid_entries = all(
            isinstance(items, list) and all(isinstance(item, dict) for item in items)
            for items in entries.values()
        )
    if not valid_entries:
        return None, f"{tool} の検査結果の項目形式が不正です"
    return data, None


def run_bandit(target: str) -> tuple:
    """bandit を実行し、Python コードの危険なパターンを検出する。
    戻り値: (findings, エラーメッセージまたはNone)
    """
    # テストコード(test_*.py等)はassert多用などでノイズが多いため対象から除外する
    exclude = "*/test_*.py,*test_*.py,*/*_test.py,*_test.py,*/tests/*,*/conftest.py"
    try:
        proc = subprocess.run(
            ["bandit", "-r", target, "-f", "json", "-q", "-x", exclude],
            capture_output=True, text=True, timeout=300,
        )
    except FileNotFoundError:
        msg = "bandit が見つかりません(pip install bandit でインストールしてください)"
        print(f"警告: {msg}")
        return [], msg
    except subprocess.TimeoutExpired:
        msg = "bandit の実行がタイムアウトしました"
        print(f"警告: {msg}")
        return [], msg
    # 1は「指摘あり」であり、実行失敗の2等とは区別する。
    data, err = _read_scan_result(proc, "bandit", "results", list, (0, 1))
    if err:
        return [], err

    findings = []
    sev_map = {"HIGH": "critical", "MEDIUM": "warning", "LOW": "info"}
    for r in data.get("results", []):
        findings.append({
            "ツール": "bandit",
            "重要度": sev_map.get(r.get("issue_severity", "LOW"), "info"),
            "ファイル": r.get("filename", ""),
            "行": r.get("line_number", ""),
            "内容": r.get("issue_text", ""),
            "ルールID": r.get("test_id", ""),
        })
    if data.get("errors"):
        return findings, "bandit で一部のファイルを検査できませんでした"
    if proc.returncode == 1 and not findings:
        return [], "bandit の終了コード1に対応する検出内容がありません"
    return findings, None


def requirements_from_pyproject(pyproject_path: Path) -> list:
    """pyproject.toml から依存ライブラリの指定を取り出し、requirements.txt 形式の行のリストにする。

    - PEP 621形式([project] の dependencies / optional-dependencies)はそのまま使う
    - Poetry形式([tool.poetry.dependencies])は、^ や ~ などpip非対応の書き方があるため
      ライブラリ名だけを使う(=最新版として検査する。実際の固定バージョンは poetry.lock 側にある)
    """
    try:
        import tomllib  # Python 3.11以降
    except ModuleNotFoundError:
        import tomli as tomllib  # Python 3.10用(requirements.txtに記載)

    data = tomllib.loads(pyproject_path.read_text(encoding="utf-8"))
    lines = []
    project = data.get("project", {}) or {}
    lines += [str(d) for d in (project.get("dependencies") or [])]
    for group in (project.get("optional-dependencies") or {}).values():
        lines += [str(d) for d in (group or [])]

    poetry = (data.get("tool", {}) or {}).get("poetry", {}) or {}
    for name, spec in (poetry.get("dependencies") or {}).items():
        if name.lower() == "python":
            continue
        if isinstance(spec, dict) and any(k in spec for k in ("git", "path", "url")):
            continue  # PyPI以外の指定は検査対象外
        lines.append(name)
    return lines


def run_pip_audit(target: str) -> tuple:
    """pip-audit を実行し、依存ライブラリの既知脆弱性(CVE)を検出する。
    requirements.txt を優先し、無ければ pyproject.toml の依存指定を検査する。
    戻り値: (findings, エラーメッセージまたはNone)
    """
    req_file = None
    for name in ("requirements.txt", "requirements-prod.txt"):
        candidate = Path(target) / name
        if candidate.exists():
            req_file = candidate
            break

    temp_dir = None
    if req_file is None:
        pyproject = Path(target) / "pyproject.toml"
        if not pyproject.exists():
            # 依存関係ファイルが無いのは「対象外」であって「失敗」ではない
            return [], None
        try:
            lines = requirements_from_pyproject(pyproject)
        except Exception as e:  # tomlの書式エラーなど
            msg = f"pyproject.toml を読み取れませんでした({e})"
            print(f"警告: {msg}")
            return [], msg
        if not lines:
            return [], None
        temp_dir = tempfile.TemporaryDirectory()
        req_file = Path(temp_dir.name) / "requirements-from-pyproject.txt"
        req_file.write_text("\n".join(lines) + "\n", encoding="utf-8")
        display_name = "pyproject.toml"
    else:
        display_name = req_file.name

    try:
        return _run_pip_audit_on(req_file, display_name)
    finally:
        if temp_dir is not None:
            temp_dir.cleanup()


def _run_pip_audit_on(req_file: Path, display_name: str) -> tuple:
    """requirements形式のファイルに pip-audit を実行し、結果を整形する。"""
    try:
        proc = subprocess.run(
            ["pip-audit", "-r", str(req_file), "-f", "json"],
            capture_output=True, text=True, timeout=300,
        )
    except FileNotFoundError:
        msg = "pip-audit が見つかりません(pip install pip-audit でインストールしてください)"
        print(f"警告: {msg}")
        return [], msg
    except subprocess.TimeoutExpired:
        msg = "pip-audit の実行がタイムアウトしました"
        print(f"警告: {msg}")
        return [], msg
    # 1は既知の脆弱性の検出時にも使われる。結果の内容も合わせて確認する。
    data, err = _read_scan_result(proc, "pip-audit", "dependencies", list, (0, 1))
    if err:
        return [], err

    findings = []
    skipped = False
    seen = set()  # pip-audit は同じIDの脆弱性情報が複数DBから重複して返ることがあるため、重複を除く
    for dep in data.get("dependencies", []):
        if dep.get("skip_reason"):
            skipped = True
            continue
        vulns = dep.get("vulns")
        if not isinstance(vulns, list) or not all(isinstance(v, dict) for v in vulns):
            return [], "pip-audit の依存ライブラリの検査結果が不正です"
        for vuln in vulns:
            key = (dep.get("name", ""), dep.get("version", ""), vuln.get("id", ""))
            if key in seen:
                continue
            seen.add(key)
            findings.append({
                "ツール": "pip-audit",
                "重要度": "critical",
                "ファイル": display_name,
                "行": "",
                "内容": f"{dep.get('name', '')} {dep.get('version', '')} に既知の脆弱性 {vuln.get('id', '')}",
                "ルールID": vuln.get("id", ""),
            })
    if skipped:
        return findings, "pip-audit で検査できなかった依存ライブラリがあります"
    if proc.returncode == 1 and not findings:
        return [], "pip-audit の終了コード1に対応する検出内容がありません"
    return findings, None


def run_detect_secrets(target: str) -> tuple:
    """detect-secrets を実行し、APIキー等のハードコードを検出する。
    戻り値: (findings, エラーメッセージまたはNone)
    """
    try:
        proc = subprocess.run(
            # --all-files を付けないと、gitで管理されていないファイルがスキップされてしまう
            ["detect-secrets", "scan", "--all-files", target],
            capture_output=True, text=True, timeout=300,
        )
    except FileNotFoundError:
        msg = "detect-secrets が見つかりません(pip install detect-secrets でインストールしてください)"
        print(f"警告: {msg}")
        return [], msg
    except subprocess.TimeoutExpired:
        msg = "detect-secrets の実行がタイムアウトしました"
        print(f"警告: {msg}")
        return [], msg
    # scanコマンドは、秘密情報が見つかった場合も正常完了なら0を返す。
    data, err = _read_scan_result(proc, "detect-secrets", "results", dict, (0,))
    if err:
        return [], err

    findings = []
    for filename, items in data.get("results", {}).items():
        for item in items:
            findings.append({
                "ツール": "detect-secrets",
                "重要度": "warning",
                "ファイル": filename,
                "行": item.get("line_number", ""),
                "内容": f"秘密情報らしき値を検出({item.get('type', '')})。誤検知の可能性もあります",
                "ルールID": item.get("type", ""),
            })
    return findings, None


def build_report(target: str, out_path: str):
    print(f"検査対象: {target}")
    errors = []

    print("bandit (コードの危険なパターン) を実行中...")
    findings, err = run_bandit(target)
    if err:
        errors.append(("bandit", err))
    print(f"  → {len(findings)} 件")

    print("pip-audit (依存ライブラリの既知脆弱性) を実行中...")
    pa, err = run_pip_audit(target)
    if err:
        errors.append(("pip-audit", err))
    findings += pa
    print(f"  → {len(pa)} 件")

    print("detect-secrets (秘密情報の直書き) を実行中...")
    ds, err = run_detect_secrets(target)
    if err:
        errors.append(("detect-secrets", err))
    findings += ds
    print(f"  → {len(ds)} 件")

    if errors:
        print("\n⚠ 以下のチェックは実行できませんでした:")
        for tool, msg in errors:
            print(f"  - {tool}: {msg}")

    wb = Workbook()
    _write_summary_sheet(wb, target, findings, errors)
    _write_findings_sheet(wb, findings, errors)
    _write_checklist_sheet(wb)
    wb.save(out_path)
    print(f"\nレポートを作成しました: {out_path}")


def _write_summary_sheet(wb, target, findings, errors=None):
    errors = errors or []
    ws = wb.active
    ws.title = "検査結果サマリー"
    ws.append(["セキュリティ公開前チェック結果"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["検査対象", target])
    ws.append(["検査日時", datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws.append([])

    counts = {"critical": 0, "warning": 0, "info": 0}
    for f in findings:
        counts[f["重要度"]] = counts.get(f["重要度"], 0) + 1

    ws.append(["重要度", "件数"])
    for cell in ws[ws.max_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for sev in ("critical", "warning", "info"):
        ws.append([SEV_LABEL_JA[sev], counts[sev]])
        ws.cell(row=ws.max_row, column=1).fill = SEV_FILL[sev]

    ws.append([])
    if errors:
        judge = "⚠ 一部のチェックが実行できませんでした(下記参照・手動で確認してください)"
    elif sum(counts.values()):
        judge = "要確認(自動チェックで検出あり)"
    else:
        judge = "自動チェックでは検出なし"
    ws.append(["総合判定", judge])

    if errors:
        ws.append([])
        ws.append(["実行できなかったチェック"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for tool, msg in errors:
            ws.append([tool, msg])
            ws.cell(row=ws.max_row, column=1).fill = SEV_FILL["critical"]

    ws.append([])
    ws.append(["※ このツールの結論はあくまで「自動チェックでの気づき」です。"])
    ws.append(["　 最終判断はご自身で行ってください。誤検知・見落としの可能性があります。"])
    ws.append(["　 自動チェックの対象は Python コードと、requirements.txt または pyproject.toml の依存指定です。"])

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 50


def _write_findings_sheet(wb, findings, errors=None):
    ws = wb.create_sheet("自動チェック結果")
    headers = ["重要度", "ツール", "ファイル", "行", "内容", "ルールID"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    if not findings:
        message = ("チェック未完了のため、指摘の有無は判断できません(検査結果サマリー参照)"
                   if errors else "自動チェックでは指摘事項は見つかりませんでした")
        ws.append(["—", "—", "—", "—", message, "—"])
    else:
        # 重要度順(critical→warning→info)に並べる
        order = {"critical": 0, "warning": 1, "info": 2}
        for f in sorted(findings, key=lambda x: order.get(x["重要度"], 9)):
            row = [SEV_LABEL_JA[f["重要度"]], f["ツール"], f["ファイル"], f["行"], f["内容"], f["ルールID"]]
            ws.append(row)
            ws.cell(row=ws.max_row, column=1).fill = SEV_FILL[f["重要度"]]

    widths = [10, 14, 30, 8, 60, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _write_checklist_sheet(wb):
    ws = wb.create_sheet("手動チェックリスト(100項目)")
    headers = ["No.", "カテゴリ", "チェック項目", "確認(OK/NG/-)", "対応のヒント", "自動チェック対象の目安"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    current_cat = None
    for item in CHECKLIST:
        cat_label = CATEGORIES[item["category"]]
        ws.append([
            item["id"],
            cat_label,
            item["item"],
            "",
            item["hint"],
            "○" if item["auto"] else "",
        ])
        row = ws.max_row
        if item["category"] != current_cat:
            # カテゴリの先頭行を少し強調
            for c in range(1, 7):
                ws.cell(row=row, column=c).font = Font(bold=True)
            current_cat = item["category"]

    widths = [6, 32, 55, 14, 45, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    parser = argparse.ArgumentParser(description="セキュリティ公開前チェックツール")
    parser.add_argument("target", help="検査したいプロジェクトのフォルダ")
    parser.add_argument("-o", "--output", default="security_report.xlsx", help="出力するExcelファイル名")
    args = parser.parse_args()

    if not Path(args.target).is_dir():
        print(f"エラー: フォルダが見つかりません: {args.target}")
        sys.exit(1)

    build_report(args.target, args.output)


if __name__ == "__main__":
    main()
