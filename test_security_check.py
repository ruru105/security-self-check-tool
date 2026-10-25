# -*- coding: utf-8 -*-
"""最低限の自動テスト。チェックリストの整合性と、レポート生成が壊れていないかを確認する。"""
from pathlib import Path

from checklist_data import CATEGORIES, CHECKLIST
from security_check import build_report


def test_checklist_has_100_items():
    assert len(CHECKLIST) == 100


def test_checklist_ids_are_unique_and_sequential():
    ids = sorted(item["id"] for item in CHECKLIST)
    assert ids == list(range(1, 101))


def test_checklist_categories_are_all_known():
    for item in CHECKLIST:
        assert item["category"] in CATEGORIES


def test_build_report_on_sample_project(tmp_path):
    out = tmp_path / "report.xlsx"
    build_report("sample_project", str(out))
    assert out.exists()

    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb.sheetnames == ["検査結果サマリー", "自動チェック結果", "手動チェックリスト(100項目)"]
    # サンプルプロジェクトには意図的な問題を仕込んであるので、何かしら検出されるはず
    findings_ws = wb["自動チェック結果"]
    assert findings_ws.max_row > 1


# ─────────────────────────────────────────────
# 以下、各ツールの結果の読み取り部分のテスト。
# 実際の bandit / pip-audit / detect-secrets は呼ばず、出力(JSON)を擬似的に用意して確認する。
# ─────────────────────────────────────────────
import json
import subprocess
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

import security_check


def _fake_run(stdout):
    """subprocess.run の代わりに、決まったJSONを返す関数を作る。"""
    def fake(cmd, **kwargs):
        return SimpleNamespace(stdout=json.dumps(stdout), returncode=1)
    return fake


def test_run_bandit_parses_severity(monkeypatch):
    out = {"results": [
        {"issue_severity": "HIGH", "filename": "a.py", "line_number": 3, "issue_text": "eval", "test_id": "B307"},
        {"issue_severity": "LOW", "filename": "b.py", "line_number": 9, "issue_text": "assert", "test_id": "B101"},
    ]}
    monkeypatch.setattr(security_check.subprocess, "run", _fake_run(out))
    findings, err = security_check.run_bandit("dummy")
    assert err is None
    assert [f["重要度"] for f in findings] == ["critical", "info"]
    assert findings[0]["ルールID"] == "B307"


@pytest.mark.parametrize("runner", ["run_bandit", "run_detect_secrets"])
def test_missing_tool_is_reported_as_error(monkeypatch, runner):
    def raise_missing(*a, **k):
        raise FileNotFoundError()
    monkeypatch.setattr(security_check.subprocess, "run", raise_missing)
    findings, err = getattr(security_check, runner)("dummy")
    assert findings == []
    assert "見つかりません" in err


def test_tool_timeout_is_reported_as_error(monkeypatch):
    def raise_timeout(*a, **k):
        raise subprocess.TimeoutExpired(cmd="bandit", timeout=300)
    monkeypatch.setattr(security_check.subprocess, "run", raise_timeout)
    findings, err = security_check.run_bandit("dummy")
    assert findings == []
    assert "タイムアウト" in err


def test_pip_audit_without_dependency_file_is_not_an_error(tmp_path):
    findings, err = security_check.run_pip_audit(str(tmp_path))
    assert findings == []
    assert err is None


def test_pip_audit_removes_duplicate_reports(monkeypatch, tmp_path):
    (tmp_path / "requirements.txt").write_text("requests==2.6.0\n", encoding="utf-8")
    out = {"dependencies": [{"name": "requests", "version": "2.6.0", "vulns": [
        {"id": "PYSEC-1"}, {"id": "PYSEC-1"}, {"id": "PYSEC-2"},
    ]}]}
    monkeypatch.setattr(security_check.subprocess, "run", _fake_run(out))
    findings, err = security_check.run_pip_audit(str(tmp_path))
    assert err is None
    assert sorted(f["ルールID"] for f in findings) == ["PYSEC-1", "PYSEC-2"]
    assert all(f["ファイル"] == "requirements.txt" for f in findings)


def test_requirements_from_pyproject_pep621_and_poetry(tmp_path):
    pyproject = tmp_path / "pyproject.toml"
    pyproject.write_text(
        '[project]\n'
        'dependencies = ["requests==2.6.0"]\n'
        '[project.optional-dependencies]\n'
        'db = ["psycopg2==2.9.9"]\n'
        '[tool.poetry.dependencies]\n'
        'python = "^3.10"\n'
        'flask = "^3.0"\n'
        'local = {path = "../local"}\n',
        encoding="utf-8",
    )
    lines = security_check.requirements_from_pyproject(pyproject)
    assert lines == ["requests==2.6.0", "psycopg2==2.9.9", "flask"]


def test_pip_audit_uses_pyproject_when_no_requirements(monkeypatch, tmp_path):
    (tmp_path / "pyproject.toml").write_text(
        '[project]\ndependencies = ["requests==2.6.0"]\n', encoding="utf-8"
    )
    seen = {}

    def fake(cmd, **kwargs):
        seen["content"] = open(cmd[cmd.index("-r") + 1], encoding="utf-8").read()
        return SimpleNamespace(stdout=json.dumps({"dependencies": [
            {"name": "requests", "version": "2.6.0", "vulns": [{"id": "PYSEC-9"}]}
        ]}), returncode=1)

    monkeypatch.setattr(security_check.subprocess, "run", fake)
    findings, err = security_check.run_pip_audit(str(tmp_path))
    assert err is None
    assert seen["content"].strip() == "requests==2.6.0"
    assert findings[0]["ファイル"] == "pyproject.toml"


def test_detect_secrets_parses_results(monkeypatch):
    out = {"results": {"app.py": [{"line_number": 5, "type": "Secret Keyword"}]}}
    monkeypatch.setattr(security_check.subprocess, "run", _fake_run(out))
    findings, err = security_check.run_detect_secrets("dummy")
    assert err is None
    assert findings[0]["重要度"] == "warning"
    assert findings[0]["行"] == 5


def test_summary_shows_warning_when_a_check_failed():
    wb = Workbook()
    security_check._write_summary_sheet(wb, "dummy", [], errors=[("bandit", "見つかりません")])
    values = [c.value for row in wb.active.iter_rows() for c in row if c.value]
    assert any("実行できませんでした" in str(v) for v in values)
    # 失敗があるのに「検出なし」と表示されないこと
    assert not any(v == "自動チェックでは検出なし" for v in values)
